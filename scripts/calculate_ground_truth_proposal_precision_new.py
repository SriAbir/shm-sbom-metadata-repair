#!/usr/bin/env python3
"""
Compare SHM semantic decisions against the human-curated ground-truth workbook.

This script answers five questions:

1. How many VERIFIED/SUPPORTED missing-field GT cases received a semantic proposal?
2. For GT-overlapping proposals, how many were accept, accept_with_caution,
   manual_review, or reject?
3. Within each decision, how many proposed values match or differ from GT?
4. What is the strict precision of automatically accepted repairs?
5. As a secondary diagnostic, what TP/FP/FN/TN result is obtained if
   semantic decision='accept' is treated as the positive prediction and
   candidate-value agreement with GT is treated as the actual class?

Important evaluation convention
-------------------------------
The legacy hash GT is not assumed to identify a unique distribution artifact.
Therefore hash_values is descriptive by default EXCEPT for cases where the
original SBOM itself identifies exactly one distribution artifact. Such cases
are exact-artifact evaluable. Use --include-hash-primary only to override this
and include all hash GT cases after a separate re-verification.

Occurrence matching is exact:
    SBOM basename + generator tool + component_index + field_name

The same package/version may appear in multiple SBOMs, so global package
identity is intentionally not used to match GT cases to proposals.

Default inputs:
    ground_truth_sample_500.xlsx
    data/full_semantic_rule_validation_results.csv

Default outputs:
    data/ground_truth_comparison_results/
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple
from urllib.parse import unquote, urlsplit, urlunsplit

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required. Install it with: pip install openpyxl"
    ) from exc


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent

TARGET_FIELDS = [
    "license",
    "supplier",
    "repository_url",
    "hash_values",
    "distribution_url",
]

PRIMARY_FIELDS_WITHOUT_HASH = {
    "license",
    "supplier",
    "repository_url",
    "distribution_url",
}

VERIFIED_GT_STATUSES = {"verified", "supported"}

DECISION_ORDER = [
    "accept",
    "accept_with_caution",
    "manual_review",
    "reject",
]

GT_COLUMN_MAP = {
    "license": ("observed_license", "gt_license", "license_status"),
    "supplier": ("observed_supplier", "gt_supplier", "supplier_status"),
    "repository_url": (
        "observed_repository_url",
        "gt_repository_url",
        "repository_url_status",
    ),
    "hash_values": (
        "observed_hash_values",
        "gt_hash_values",
        "hash_values_status",
    ),
    "distribution_url": (
        "observed_distribution_url",
        "gt_distribution_url",
        "distribution_url_status",
    ),
}

PROJECT_ID_RE = re.compile(r"(?:^|/)(NPM|PY|MAV)-\d+", re.IGNORECASE)

LICENSE_ALIASES = {
    "apache license 2.0": "apache-2.0",
    "apache 2.0": "apache-2.0",
    "apache software license": "apache-2.0",
    "mit license": "mit",
    "bsd-3-clause license": "bsd-3-clause",
    "bsd 3-clause license": "bsd-3-clause",
    "bsd-2-clause license": "bsd-2-clause",
    "bsd 2-clause license": "bsd-2-clause",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def lower(value) -> str:
    return clean(value).lower()


def is_blank(value) -> bool:
    text = clean(value)
    return text == "" or text.lower() in {
        "null", "none", "unknown", "n/a", "na", "[]", "{}"
    }


def pct(num: int, den: int) -> float:
    return (100.0 * num / den) if den else 0.0


def safe_div(num: float, den: float) -> float:
    return (num / den) if den else 0.0


def normalize_space(value) -> str:
    return re.sub(r"\s+", " ", clean(value)).strip()


def normalize_name(value) -> str:
    return normalize_space(value).casefold()


def normalize_version(value) -> str:
    return normalize_space(value).casefold()


def normalize_purl(value) -> str:
    text = clean(value)
    if not text:
        return ""
    return unquote(text).strip().casefold()


def normalize_component_index(value) -> str:
    text = clean(value)
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def canonical_url(value) -> str:
    """Normalize harmless representation differences in a single URL."""
    value = clean(value)
    if not value:
        return ""

    if value.lower().startswith("scm:git:"):
        value = value[len("scm:git:"):]
    if value.lower().startswith("git+"):
        value = value[4:]
    if value.lower().startswith("git://"):
        value = "https://" + value[len("git://"):]

    if value.lower().startswith("ssh://git@"):
        value = "https://" + value[len("ssh://git@"):]
    elif value.lower().startswith("git@") and ":" in value:
        host_path = value[len("git@"):]
        host, path = host_path.split(":", 1)
        value = f"https://{host}/{path}"

    try:
        parsed = urlsplit(value)
        if not parsed.netloc:
            return value.rstrip("/").casefold()

        scheme = parsed.scheme.lower() or "https"
        host = parsed.netloc.lower()
        path = unquote(parsed.path).rstrip("/")
        if path.lower().endswith(".git"):
            path = path[:-4]

        return urlunsplit((scheme, host, path, "", "")).casefold()
    except Exception:
        return value.rstrip("/").casefold()


def split_multi_values(value) -> List[str]:
    return [
        part.strip()
        for part in re.split(r"[;|\n]+", clean(value))
        if part.strip()
    ]


def normalize_url_set(value) -> Tuple[str, ...]:
    values = [canonical_url(part) for part in split_multi_values(value)]
    return tuple(sorted(set(part for part in values if part)))


def normalize_license(value) -> str:
    text = normalize_space(value).casefold()
    text = re.sub(r"\s*([()])\s*", r"\1", text)
    text = re.sub(r"\s+", " ", text)
    return LICENSE_ALIASES.get(text, text)


def normalize_supplier(value) -> str:
    text = normalize_space(value)
    text = re.sub(r"^organization\s*:\s*", "", text, flags=re.IGNORECASE)
    return text.casefold()


def canonical_hash_algorithm(value) -> str:
    """Normalize common digest algorithm spellings for GT/proposal comparison."""
    text = clean(value).casefold().replace(" ", "")
    compact = re.sub(r"[-_]", "", text)
    aliases = {
        "sha1": "sha1",
        "sha224": "sha224",
        "sha256": "sha256",
        "sha384": "sha384",
        "sha512": "sha512",
        "md5": "md5",
        "blake2b256": "blake2b_256",
    }
    return aliases.get(compact, text)


def normalize_hash_payload(value) -> str:
    return re.sub(r"\s+", "", clean(value)).casefold()


def normalize_hash_token(token) -> str:
    text = clean(token)
    if not text:
        return ""

    if ":" in text:
        alg, payload = text.split(":", 1)
        if re.fullmatch(r"[A-Za-z0-9_-]+", alg.strip()):
            return f"{canonical_hash_algorithm(alg)}:{normalize_hash_payload(payload)}"

    return normalize_hash_payload(text)


def normalize_hash_set(value) -> Tuple[str, ...]:
    tokens = []
    for token in split_multi_values(value):
        normalized = normalize_hash_token(token)
        if normalized:
            tokens.append(normalized)
    return tuple(sorted(set(tokens)))


def parse_hash_entries(value, algorithms="") -> set[Tuple[str, str]]:
    """
    Parse hashes as (algorithm, digest) pairs.

    SHM's proposal value may contain only digest strings while the associated
    algorithms are preserved separately in evidence_hash_algorithms. GT values
    may instead use labels such as ``SHA-256:<digest>``.
    """
    raw_values = split_multi_values(value)
    raw_algorithms = split_multi_values(algorithms)
    entries: set[Tuple[str, str]] = set()

    for index, token in enumerate(raw_values):
        text = clean(token)
        if not text:
            continue

        if ":" in text:
            maybe_alg, payload = text.split(":", 1)
            if re.fullmatch(r"[A-Za-z0-9_-]+", maybe_alg.strip()):
                entries.add((
                    canonical_hash_algorithm(maybe_alg),
                    normalize_hash_payload(payload),
                ))
                continue

        algorithm = ""
        if len(raw_algorithms) == len(raw_values):
            algorithm = canonical_hash_algorithm(raw_algorithms[index])
        entries.add((algorithm, normalize_hash_payload(text)))

    return {(alg, digest) for alg, digest in entries if digest}


def hash_gt_supported_by_proposal(proposal_value, ground_truth_value, proposal_algorithms="") -> bool:
    """
    Return True when every GT digest is present in SHM's hashes for the SAME
    exact artifact. Extra valid algorithms in SHM do not make a repair wrong.

    Example: GT contains only SHA-256 while SHM preserves BLAKE2b-256, MD5 and
    SHA-256 for the same artifact. This is a correct match when the GT SHA-256
    digest is present in SHM's proposal.
    """
    proposal_entries = parse_hash_entries(proposal_value, proposal_algorithms)
    gt_entries = parse_hash_entries(ground_truth_value)
    if not proposal_entries or not gt_entries:
        return False

    proposal_digests = {digest for _, digest in proposal_entries}
    for gt_alg, gt_digest in gt_entries:
        if gt_alg:
            if (gt_alg, gt_digest) not in proposal_entries:
                return False
        elif gt_digest not in proposal_digests:
            return False
    return True


def normalized_value(field: str, value):
    if field == "license":
        return normalize_license(value)
    if field == "supplier":
        return normalize_supplier(value)
    if field == "repository_url":
        return canonical_url(value)
    if field == "distribution_url":
        return normalize_url_set(value)
    if field == "hash_values":
        return normalize_hash_set(value)
    if field == "purl":
        return normalize_purl(value)
    if field == "version":
        return normalize_version(value)
    return normalize_space(value).casefold()


def values_equal(field: str, proposed, ground_truth) -> bool:
    return normalized_value(field, proposed) == normalized_value(field, ground_truth)


def distribution_overlap(proposed, ground_truth) -> bool:
    proposed_set = set(normalize_url_set(proposed))
    gt_set = set(normalize_url_set(ground_truth))
    return bool(proposed_set and gt_set and proposed_set.intersection(gt_set))


def extract_project_id_from_source(source) -> str:
    text = clean(source).replace("\\", "/")
    match = PROJECT_ID_RE.search("/" + text.lstrip("/"))
    if not match:
        return ""

    prefix = match.group(1).upper()
    full = re.search(rf"{prefix}-\d+", text, re.IGNORECASE)
    return full.group(0).upper() if full else ""


def extract_tool_from_path(source) -> str:
    parts = [part.casefold() for part in Path(clean(source)).parts]
    if "syft" in parts:
        return "syft"
    if "cdxgen" in parts:
        return "cdxgen"
    return ""


def sbom_basename(source) -> str:
    return Path(clean(source)).name.casefold()


def truthy(value) -> bool:
    return lower(value) in {"true", "1", "yes", "y"}


# ---------------------------------------------------------------------------
# Input discovery
# ---------------------------------------------------------------------------

def locate_ground_truth(explicit: Path | None) -> Path:
    if explicit is not None:
        if not explicit.exists():
            raise FileNotFoundError(f"Ground-truth workbook not found: {explicit}")
        return explicit

    candidates = [
        BASE_DIR / "ground_truth_sample_500.xlsx",
        BASE_DIR / "Test" / "ground_truth_sample_500.xlsx",
        BASE_DIR / "test" / "ground_truth_sample_500.xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError(
        "Could not find ground_truth_sample_500.xlsx. Tried:\n  "
        + "\n  ".join(str(path) for path in candidates)
    )


def locate_semantic_csv(explicit: Path | None) -> Path:
    if explicit is not None:
        if not explicit.exists():
            raise FileNotFoundError(f"Semantic-validation CSV not found: {explicit}")
        return explicit

    preferred = [
        BASE_DIR / "data" / "full_semantic_rule_validation_results.csv",
        BASE_DIR / "full_semantic_rule_validation_results.csv",
        BASE_DIR / "Test" / "full_semantic_rule_validation_results.csv",
        BASE_DIR / "test" / "full_semantic_rule_validation_results.csv",
    ]
    for path in preferred:
        if path.exists():
            return path

    matches = []
    for directory in [BASE_DIR / "data", BASE_DIR / "Test", BASE_DIR / "test", BASE_DIR]:
        if directory.exists():
            matches.extend(directory.glob("full_semantic_rule_validation_results*.csv"))

    if not matches:
        raise FileNotFoundError(
            "Could not find full_semantic_rule_validation_results.csv"
        )

    matches = sorted(set(matches), key=lambda p: p.stat().st_mtime, reverse=True)
    print("WARNING: exact current semantic filename was not found.")
    print(f"Using most recently modified candidate: {matches[0]}")
    return matches[0]


# ---------------------------------------------------------------------------
# Read ground truth
# ---------------------------------------------------------------------------

def load_ground_truth(path: Path) -> Tuple[List[Dict[str, str]], List[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Ground_Truth" not in wb.sheetnames:
        raise ValueError("Workbook does not contain a 'Ground_Truth' sheet")

    ws = wb["Ground_Truth"]
    rows = ws.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]

    required = {"gt_id", "ecosystem", "name", "source_sbom", "component_index"}
    for observed_col, gt_col, status_col in GT_COLUMN_MAP.values():
        required.update({observed_col, gt_col, status_col})

    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(
            "Ground-truth workbook is missing required columns: "
            + ", ".join(missing)
        )

    output = []
    for values in rows:
        row = dict(zip(headers, values))
        if not clean(row.get("gt_id")):
            continue
        row["_project_id"] = extract_project_id_from_source(row.get("source_sbom"))
        output.append(row)

    return output, headers


def build_gt_cases(gt_rows: Sequence[Dict[str, str]]) -> Tuple[List[Dict[str, str]], int]:
    """Return VERIFIED/SUPPORTED originally-missing cases in SHM repair scope."""
    cases = []
    all_originally_missing = 0

    for row in gt_rows:
        for field in TARGET_FIELDS:
            observed_col, gt_col, status_col = GT_COLUMN_MAP[field]
            observed = clean(row.get(observed_col))
            gt_value = clean(row.get(gt_col))
            status = lower(row.get(status_col))

            if is_blank(observed):
                all_originally_missing += 1
            else:
                continue

            if status not in VERIFIED_GT_STATUSES or is_blank(gt_value):
                continue

            observed_purl = clean(row.get("observed_purl"))
            gt_purl = clean(row.get("gt_purl"))
            observed_version = clean(row.get("observed_version"))
            gt_version = clean(row.get("gt_version"))

            cases.append({
                "gt_id": clean(row.get("gt_id")),
                "ecosystem": lower(row.get("ecosystem")),
                "project_id": clean(row.get("_project_id")),
                "name": clean(row.get("name")),
                "version_for_identity": observed_version or gt_version,
                "purl_for_identity": observed_purl or gt_purl,
                "field_name": field,
                "ground_truth_value": gt_value,
                "ground_truth_status": status,
                "source_sbom": clean(row.get("source_sbom")),
                "component_index": normalize_component_index(row.get("component_index")),
                "gt_tool": extract_tool_from_path(row.get("source_sbom")),
                # Hash artifact identity must come from the original SBOM-side metadata,
                # not from GT alone, to avoid evaluation leakage.
                "observed_distribution_url": clean(row.get("observed_distribution_url")),
                "gt_distribution_url": clean(row.get("gt_distribution_url")),
            })

    return cases, all_originally_missing


# ---------------------------------------------------------------------------
# Read semantic results
# ---------------------------------------------------------------------------

def load_semantic_results(path: Path) -> List[Dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        headers = set(reader.fieldnames or [])

        required = {
            "component_id",
            "project_id",
            "name",
            "field_name",
            "new_value",
            "semantic_rule_label",
            "semantic_rule_decision",
            "tool",
            "sbom_path",
            "component_index",
        }

        missing = sorted(required - headers)
        if missing:
            raise ValueError(
                "Semantic-validation CSV is missing required columns: "
                + ", ".join(missing)
            )

        rows = [dict(row) for row in reader]

    return rows


# ---------------------------------------------------------------------------
# Exact occurrence matching
# ---------------------------------------------------------------------------

def occurrence_key_from_case(case: Dict[str, str]) -> Tuple[str, str, str, str]:
    return (
        sbom_basename(case.get("source_sbom")),
        lower(case.get("gt_tool")),
        normalize_component_index(case.get("component_index")),
        lower(case.get("field_name")),
    )


def occurrence_key_from_proposal(proposal: Dict[str, str]) -> Tuple[str, str, str, str]:
    tool = lower(proposal.get("tool")) or extract_tool_from_path(proposal.get("sbom_path"))
    return (
        sbom_basename(proposal.get("sbom_path")),
        tool,
        normalize_component_index(proposal.get("component_index")),
        lower(proposal.get("field_name")),
    )


def determine_hash_artifact_alignment(case: Dict[str, str], proposal: Dict[str, str]) -> str:
    """
    Establish exact hash-artifact alignment without using GT to select an artifact.

    By default a hash case is exact-artifact evaluable only when the ORIGINAL
    SBOM metadata identifies exactly one distribution URL and SHM's preserved
    artifact URL is that same URL. Multiple original distribution URLs remain
    ambiguous even if the proposal happens to match one of them.
    """
    if case.get("field_name") != "hash_values":
        return "not_applicable"

    proposal_artifact = clean(proposal.get("evidence_artifact_url"))
    observed_distribution = clean(case.get("observed_distribution_url"))

    if not proposal_artifact or not observed_distribution:
        return "not_established"

    observed_urls = set(normalize_url_set(observed_distribution))
    if len(observed_urls) > 1:
        return "original_sbom_has_multiple_artifacts"
    if len(observed_urls) == 0:
        return "not_established"

    artifact_norm = canonical_url(proposal_artifact)
    if artifact_norm in observed_urls:
        return "exact_artifact_url_matches_original_sbom"

    return "artifact_url_differs_from_original_sbom"


def case_primary_evaluable(case: Dict[str, str], include_hash_primary: bool) -> bool:
    """Return whether a VERIFIED GT case belongs in the primary denominator."""
    if case.get("field_name") != "hash_values":
        return True
    if include_hash_primary:
        return True

    # Without the override, admit only hash GT cases whose original SBOM
    # identifies exactly one artifact. In the current GT this admits jieba.
    observed_urls = set(normalize_url_set(case.get("observed_distribution_url")))
    return len(observed_urls) == 1


def match_cases_to_proposals(
    cases: Sequence[Dict[str, str]],
    proposals: Sequence[Dict[str, str]],
    include_hash_primary: bool,
) -> Tuple[List[Dict[str, str]], List[Dict[str, str]], int]:
    proposals_by_occurrence: Dict[Tuple[str, str, str, str], List[Tuple[int, Dict[str, str]]]] = defaultdict(list)

    for index, proposal in enumerate(proposals):
        proposals_by_occurrence[occurrence_key_from_proposal(proposal)].append((index, proposal))

    duplicate_keys = {
        key: values
        for key, values in proposals_by_occurrence.items()
        if len(values) > 1
    }

    # A component occurrence + field should have one semantic decision. Do not
    # silently inflate evaluation counts if a pipeline error introduced duplicates.
    for key, values in duplicate_keys.items():
        signatures = {
            (
                lower(row.get("semantic_rule_decision")),
                repr(normalized_value(lower(row.get("field_name")), row.get("new_value"))),
            )
            for _, row in values
        }
        if len(signatures) > 1:
            raise ValueError(
                "Conflicting duplicate semantic proposals found for occurrence key "
                f"{key}. Fix the semantic result file before GT evaluation."
            )

    matched = []
    unmatched = []

    for case in cases:
        key = occurrence_key_from_case(case)
        candidates = proposals_by_occurrence.get(key, [])

        if not candidates:
            unmatched.append(case)
            continue

        # Identical duplicates, if any, count only once.
        proposal_index, proposal = candidates[0]
        field = case["field_name"]
        proposed_value = clean(proposal.get("new_value"))
        hash_alignment = determine_hash_artifact_alignment(case, proposal)

        if field == "hash_values":
            is_correct = (
                hash_alignment == "exact_artifact_url_matches_original_sbom"
                and hash_gt_supported_by_proposal(
                    proposed_value,
                    case["ground_truth_value"],
                    proposal.get("evidence_hash_algorithms", ""),
                )
            )
        else:
            is_correct = values_equal(field, proposed_value, case["ground_truth_value"])

        gt_project = clean(case.get("project_id")).upper()
        proposal_project = clean(proposal.get("project_id")).upper()
        project_consistent = (
            not gt_project or not proposal_project or gt_project == proposal_project
        )
        name_consistent = (
            not normalize_name(case.get("name"))
            or normalize_name(case.get("name")) == normalize_name(proposal.get("name"))
        )

        if field != "hash_values":
            primary_evaluable = True
            primary_exclusion_reason = ""
        elif include_hash_primary:
            primary_evaluable = True
            primary_exclusion_reason = ""
        elif hash_alignment == "exact_artifact_url_matches_original_sbom":
            primary_evaluable = True
            primary_exclusion_reason = ""
        else:
            primary_evaluable = False
            primary_exclusion_reason = "hash_artifact_identity_not_established_from_original_sbom"

        distribution_partial_overlap = ""
        if field == "distribution_url" and not is_correct:
            distribution_partial_overlap = str(
                distribution_overlap(proposed_value, case["ground_truth_value"])
            ).lower()

        matched.append({
            "gt_id": case["gt_id"],
            "field_name": field,
            "ecosystem": case["ecosystem"],
            "gt_project_id": case["project_id"],
            "gt_tool": case["gt_tool"],
            "gt_source_sbom": case["source_sbom"],
            "gt_component_index": case["component_index"],
            "name": case["name"],
            "identity_version": case["version_for_identity"],
            "identity_purl": case["purl_for_identity"],
            "ground_truth_status": case["ground_truth_status"],
            "ground_truth_value": case["ground_truth_value"],
            "proposal_value": proposed_value,
            "value_match": "correct" if is_correct else "incorrect",
            "distribution_partial_overlap": distribution_partial_overlap,
            "semantic_rule_label": lower(proposal.get("semantic_rule_label")),
            "semantic_rule_decision": lower(proposal.get("semantic_rule_decision")),
            "semantic_rule_id": clean(proposal.get("semantic_rule_id")),
            "semantic_rule_reason": clean(proposal.get("semantic_rule_reason")),
            "static_validation_result": lower(proposal.get("static_validation_result")),
            "proposal_component_id": clean(proposal.get("component_id")),
            "proposal_project_id": clean(proposal.get("project_id")),
            "proposal_tool": clean(proposal.get("tool")),
            "proposal_sbom_path": clean(proposal.get("sbom_path")),
            "proposal_component_index": normalize_component_index(proposal.get("component_index")),
            "evidence_party_role": clean(proposal.get("evidence_party_role")),
            "evidence_party_source_field": clean(proposal.get("evidence_party_source_field")),
            "evidence_artifact_url": clean(proposal.get("evidence_artifact_url")),
            "evidence_artifact_match": clean(proposal.get("evidence_artifact_match")),
            "evidence_artifact_match_method": clean(proposal.get("evidence_artifact_match_method")),
            "evidence_hash_usable_for_auto_repair": clean(proposal.get("evidence_hash_usable_for_auto_repair")),
            "observed_distribution_url": case.get("observed_distribution_url", ""),
            "gt_distribution_url": case.get("gt_distribution_url", ""),
            "hash_artifact_alignment": hash_alignment,
            "hash_comparison_mode": (
                "gt_hash_subset_of_same_artifact_proposal"
                if field == "hash_values" and hash_alignment == "exact_artifact_url_matches_original_sbom"
                else "not_applicable"
            ),
            "primary_evaluable": "yes" if primary_evaluable else "no",
            "primary_exclusion_reason": primary_exclusion_reason,
            "project_consistent": str(project_consistent).lower(),
            "name_consistent": str(name_consistent).lower(),
            "proposal_row_number": proposal_index + 2,
            "duplicate_proposal_rows_for_key": len(candidates),
        })

    return matched, unmatched, len(duplicate_keys)


# ---------------------------------------------------------------------------
# Metrics and summaries
# ---------------------------------------------------------------------------

def primary_records(records: Sequence[Dict[str, str]]) -> List[Dict[str, str]]:
    return [row for row in records if row.get("primary_evaluable") == "yes"]


def decision_summary(
    matched: Sequence[Dict[str, str]],
    unmatched: Sequence[Dict[str, str]],
) -> List[Dict[str, str]]:
    total_gt = len(matched) + len(unmatched)
    rows = []

    for decision in DECISION_ORDER:
        subset = [row for row in matched if row["semantic_rule_decision"] == decision]
        correct = sum(row["value_match"] == "correct" for row in subset)
        incorrect = sum(row["value_match"] == "incorrect" for row in subset)
        primary_count = sum(row["primary_evaluable"] == "yes" for row in subset)

        rows.append({
            "semantic_rule_decision": decision,
            "verified_gt_overlap": len(subset),
            "candidate_matches_gt": correct,
            "candidate_differs_from_gt": incorrect,
            "candidate_match_rate_pct": f"{pct(correct, len(subset)):.2f}",
            "primary_evaluable_overlap": primary_count,
            "percentage_of_verified_gt_cases": f"{pct(len(subset), total_gt):.2f}",
        })

    rows.append({
        "semantic_rule_decision": "no_proposal",
        "verified_gt_overlap": len(unmatched),
        "candidate_matches_gt": "",
        "candidate_differs_from_gt": "",
        "candidate_match_rate_pct": "",
        "primary_evaluable_overlap": "",
        "percentage_of_verified_gt_cases": f"{pct(len(unmatched), total_gt):.2f}",
    })

    return rows


def decision_summary_by_field(
    matched: Sequence[Dict[str, str]],
    unmatched: Sequence[Dict[str, str]],
) -> List[Dict[str, str]]:
    rows = []

    for field in TARGET_FIELDS:
        field_matched = [row for row in matched if row["field_name"] == field]
        field_unmatched = [row for row in unmatched if row["field_name"] == field]
        total = len(field_matched) + len(field_unmatched)

        for decision in DECISION_ORDER:
            subset = [
                row for row in field_matched
                if row["semantic_rule_decision"] == decision
            ]
            if not subset:
                continue

            correct = sum(row["value_match"] == "correct" for row in subset)
            incorrect = sum(row["value_match"] == "incorrect" for row in subset)

            rows.append({
                "field_name": field,
                "semantic_rule_decision": decision,
                "verified_gt_overlap": len(subset),
                "candidate_matches_gt": correct,
                "candidate_differs_from_gt": incorrect,
                "candidate_match_rate_pct": f"{pct(correct, len(subset)):.2f}",
                "percentage_of_field_verified_gt_cases": f"{pct(len(subset), total):.2f}",
            })

        if field_unmatched:
            rows.append({
                "field_name": field,
                "semantic_rule_decision": "no_proposal",
                "verified_gt_overlap": len(field_unmatched),
                "candidate_matches_gt": "",
                "candidate_differs_from_gt": "",
                "candidate_match_rate_pct": "",
                "percentage_of_field_verified_gt_cases": f"{pct(len(field_unmatched), total):.2f}",
            })

    return rows


def decision_value_matrix(records: Sequence[Dict[str, str]]) -> List[Dict[str, str]]:
    counts = Counter(
        (row["semantic_rule_decision"], row["value_match"])
        for row in records
    )

    rows = []
    for decision in DECISION_ORDER:
        for outcome in ["correct", "incorrect"]:
            rows.append({
                "semantic_rule_decision": decision,
                "ground_truth_value_match": outcome,
                "count": counts.get((decision, outcome), 0),
            })
    return rows


def strict_precision(records: Sequence[Dict[str, str]], positive_decisions: set[str]) -> Dict[str, float]:
    selected = [
        row for row in primary_records(records)
        if row["semantic_rule_decision"] in positive_decisions
    ]
    correct = sum(row["value_match"] == "correct" for row in selected)
    incorrect = len(selected) - correct

    return {
        "evaluated": len(selected),
        "correct": correct,
        "incorrect": incorrect,
        "precision_pct": pct(correct, len(selected)),
    }


def precision_by_field(records: Sequence[Dict[str, str]], positive_decisions: set[str]) -> List[Dict[str, str]]:
    rows = []
    evaluable = primary_records(records)

    for field in TARGET_FIELDS:
        subset = [
            row for row in evaluable
            if row["field_name"] == field
            and row["semantic_rule_decision"] in positive_decisions
        ]
        correct = sum(row["value_match"] == "correct" for row in subset)
        rows.append({
            "field_name": field,
            "evaluated_accepted_proposals": len(subset),
            "correct": correct,
            "incorrect": len(subset) - correct,
            "precision_pct": f"{pct(correct, len(subset)):.2f}",
        })

    return rows


def binary_confusion_diagnostic(
    records: Sequence[Dict[str, str]],
    positive_decisions: set[str],
    label: str,
) -> Dict[str, str]:
    """
    Secondary diagnostic over GT cases that HAVE a proposal.

    Predicted positive: semantic decision is in positive_decisions.
    Actual positive: candidate value matches independently verified GT.

    This is not an end-to-end confusion matrix because no-proposal GT cases
    have no candidate value and therefore cannot be assigned candidate-correct
    versus candidate-incorrect status.
    """
    rows = primary_records(records)

    tp = fp = fn = tn = 0
    for row in rows:
        predicted_positive = row["semantic_rule_decision"] in positive_decisions
        actual_positive = row["value_match"] == "correct"

        if predicted_positive and actual_positive:
            tp += 1
        elif predicted_positive and not actual_positive:
            fp += 1
        elif not predicted_positive and actual_positive:
            fn += 1
        else:
            tn += 1

    precision = safe_div(tp, tp + fp)
    recall = safe_div(tp, tp + fn)
    f1 = safe_div(2 * precision * recall, precision + recall)
    specificity = safe_div(tn, tn + fp)
    accuracy = safe_div(tp + tn, tp + fp + fn + tn)

    return {
        "analysis": label,
        "population": "primary-evaluable VERIFIED GT cases with a semantic proposal",
        "tp": str(tp),
        "fp": str(fp),
        "fn": str(fn),
        "tn": str(tn),
        "precision_pct": f"{precision * 100:.2f}",
        "recall_pct": f"{recall * 100:.2f}",
        "f1_pct": f"{f1 * 100:.2f}",
        "specificity_pct": f"{specificity * 100:.2f}",
        "accuracy_pct": f"{accuracy * 100:.2f}",
    }


def end_to_end_metrics(
    all_gt_cases: Sequence[Dict[str, str]],
    matched: Sequence[Dict[str, str]],
    include_hash_primary: bool,
) -> Dict[str, str]:
    denominator_cases = [
        case for case in all_gt_cases
        if case_primary_evaluable(case, include_hash_primary)
    ]

    evaluable_matched = primary_records(matched)
    accepted = [
        row for row in evaluable_matched
        if row["semantic_rule_decision"] == "accept"
    ]
    correct_accepted = sum(row["value_match"] == "correct" for row in accepted)
    incorrect_accepted = sum(row["value_match"] == "incorrect" for row in accepted)

    precision = safe_div(correct_accepted, correct_accepted + incorrect_accepted)
    recall = safe_div(correct_accepted, len(denominator_cases))
    f1 = safe_div(2 * precision * recall, precision + recall)

    covered_keys = {
        (row["gt_id"], row["field_name"])
        for row in evaluable_matched
    }
    denominator_keys = {
        (case["gt_id"], case["field_name"])
        for case in denominator_cases
    }

    return {
        "primary_verified_gt_cases": str(len(denominator_cases)),
        "primary_gt_cases_with_proposal": str(len(covered_keys & denominator_keys)),
        "primary_proposal_coverage_pct": f"{pct(len(covered_keys & denominator_keys), len(denominator_cases)):.2f}",
        "strict_accepted_evaluated": str(len(accepted)),
        "strict_correct_accepted": str(correct_accepted),
        "strict_incorrect_accepted": str(incorrect_accepted),
        "strict_precision_pct": f"{precision * 100:.2f}",
        "verified_repair_recall_pct": f"{recall * 100:.2f}",
        "end_to_end_f1_pct": f"{f1 * 100:.2f}",
    }


def field_recovery_summary(
    gt_cases: Sequence[Dict[str, str]],
    matched: Sequence[Dict[str, str]],
    include_hash_primary: bool,
) -> List[Dict[str, str]]:
    rows = []

    for field in TARGET_FIELDS:
        field_gt = [case for case in gt_cases if case["field_name"] == field]
        field_matched = [row for row in matched if row["field_name"] == field]
        primary_gt = [
            case for case in field_gt
            if case_primary_evaluable(case, include_hash_primary)
        ]
        evaluable_matched = [
            row for row in field_matched
            if row["primary_evaluable"] == "yes"
        ]

        accepted = [
            row for row in evaluable_matched
            if row["semantic_rule_decision"] == "accept"
        ]
        correct = sum(row["value_match"] == "correct" for row in accepted)
        incorrect = sum(row["value_match"] == "incorrect" for row in accepted)

        if field != "hash_values":
            primary_label = "yes"
        elif include_hash_primary:
            primary_label = "yes_all_hash_cases"
        else:
            primary_label = "exact_artifact_only"

        rows.append({
            "field_name": field,
            "verified_gt_cases": len(field_gt),
            "gt_cases_with_any_proposal": len(field_matched),
            "proposal_coverage_pct": f"{pct(len(field_matched), len(field_gt)):.2f}",
            "primary_metric_included": primary_label,
            "primary_evaluable_gt_cases": len(primary_gt),
            "primary_evaluable_cases_with_proposal": len(evaluable_matched),
            "strict_accepted_evaluated": len(accepted),
            "correct_accepted": correct,
            "incorrect_accepted": incorrect,
            "strict_precision_pct": f"{pct(correct, len(accepted)):.2f}",
            "verified_recovery_pct": f"{pct(correct, len(primary_gt)):.2f}" if primary_gt else "",
        })

    return rows


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

def write_csv(path: Path, rows: Iterable[Dict[str, object]], fieldnames: List[str] | None = None) -> None:
    rows = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)

    if fieldnames is None:
        fieldnames = list(rows[0].keys()) if rows else []

    with path.open("w", newline="", encoding="utf-8") as handle:
        if not fieldnames:
            return
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def print_decision_table(rows: Sequence[Dict[str, object]]) -> None:
    print("\nVERIFIED GT CASES BY SEMANTIC DECISION")
    print("--------------------------------------")
    for row in rows:
        decision = str(row["semantic_rule_decision"])
        overlap = row["verified_gt_overlap"]
        correct = row["candidate_matches_gt"]
        incorrect = row["candidate_differs_from_gt"]

        if decision == "no_proposal":
            print(f"{decision:22} GT cases={overlap}")
        else:
            print(
                f"{decision:22} GT overlap={overlap:4}  "
                f"match={correct:4}  mismatch={incorrect:4}"
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare SHM semantic decisions with VERIFIED ground truth."
    )
    parser.add_argument(
        "--ground-truth",
        type=Path,
        default=None,
        help="Path to ground_truth_sample_500.xlsx",
    )
    parser.add_argument(
        "--semantic",
        type=Path,
        default=None,
        help="Path to full_semantic_rule_validation_results.csv",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=BASE_DIR / "data" / "ground_truth_comparison_results",
    )
    parser.add_argument(
        "--include-hash-primary",
        action="store_true",
        help=(
            "Include hash_values in primary precision/recall metrics. "
            "Use only after GT hashes have exact-artifact identity verification."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    gt_path = locate_ground_truth(args.ground_truth)
    semantic_path = locate_semantic_csv(args.semantic)
    output_dir = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 78)
    print("SHM Ground-Truth Decision and Repair Evaluation")
    print("=" * 78)
    print(f"Ground truth:       {gt_path}")
    print(f"Semantic results:   {semantic_path}")
    print(f"Output directory:   {output_dir}")
    print()

    gt_rows, _ = load_ground_truth(gt_path)
    semantic_rows = load_semantic_results(semantic_path)
    gt_cases, all_missing = build_gt_cases(gt_rows)

    matched, unmatched, duplicate_key_count = match_cases_to_proposals(
        gt_cases,
        semantic_rows,
        include_hash_primary=args.include_hash_primary,
    )

    status_counts = Counter(case["ground_truth_status"] for case in gt_cases)
    gt_field_counts = Counter(case["field_name"] for case in gt_cases)

    print(f"Ground-truth component occurrences:       {len(gt_rows)}")
    print(f"Originally missing target-field cases:    {all_missing}")
    print(f"VERIFIED/SUPPORTED GT missing cases:      {len(gt_cases)}")
    print(f"Semantic result rows:                     {len(semantic_rows)}")
    print(f"GT cases with semantic proposal:          {len(matched)}")
    print(f"GT cases without semantic proposal:       {len(unmatched)}")
    print(f"Duplicate semantic occurrence keys:       {duplicate_key_count}")

    print("\nGT status counts:")
    for status, count in sorted(status_counts.items()):
        print(f"  {status}: {count}")

    print("\nVERIFIED/SUPPORTED GT cases by field:")
    for field in TARGET_FIELDS:
        print(f"  {field:20} {gt_field_counts.get(field, 0)}")

    decisions = decision_summary(matched, unmatched)
    print_decision_table(decisions)

    strict = strict_precision(matched, {"accept"})
    expanded = strict_precision(matched, {"accept", "accept_with_caution"})
    e2e = end_to_end_metrics(gt_cases, matched, args.include_hash_primary)

    print("\nPRIMARY: STRICT AUTOMATIC ACCEPTANCE")
    print("------------------------------------")
    print(f"Accepted GT-overlap proposals evaluated: {strict['evaluated']}")
    print(f"Correct accepted proposals:              {strict['correct']}")
    print(f"Incorrect accepted proposals:            {strict['incorrect']}")
    print(f"Strict precision:                        {strict['precision_pct']:.2f}%")
    print(f"Verified repair recall:                  {e2e['verified_repair_recall_pct']}%")
    print(f"End-to-end F1:                           {e2e['end_to_end_f1_pct']}%")
    print(f"Primary proposal coverage:               {e2e['primary_proposal_coverage_pct']}%")

    print("\nSECONDARY: ACCEPT + ACCEPT_WITH_CAUTION")
    print("----------------------------------------")
    print(f"Evaluated:                                {expanded['evaluated']}")
    print(f"Correct:                                  {expanded['correct']}")
    print(f"Incorrect:                                {expanded['incorrect']}")
    print(f"Precision:                                {expanded['precision_pct']:.2f}%")

    strict_confusion = binary_confusion_diagnostic(
        matched,
        {"accept"},
        "strict_accept_vs_non_accept",
    )
    expanded_confusion = binary_confusion_diagnostic(
        matched,
        {"accept", "accept_with_caution"},
        "accept_or_caution_vs_manual_or_reject",
    )

    print("\nSECONDARY DIAGNOSTIC: TP / FP / FN / TN")
    print("-----------------------------------------")
    print(
        "Population: primary-evaluable VERIFIED GT cases that received a proposal"
    )
    print(
        f"Strict accept: TP={strict_confusion['tp']} FP={strict_confusion['fp']} "
        f"FN={strict_confusion['fn']} TN={strict_confusion['tn']}"
    )
    print(
        f"Diagnostic recall={strict_confusion['recall_pct']}%  "
        f"F1={strict_confusion['f1_pct']}%"
    )

    # Detail outputs.
    detail_fields = [
        "gt_id",
        "field_name",
        "ecosystem",
        "gt_project_id",
        "gt_tool",
        "gt_source_sbom",
        "gt_component_index",
        "name",
        "identity_version",
        "identity_purl",
        "ground_truth_status",
        "ground_truth_value",
        "proposal_value",
        "value_match",
        "distribution_partial_overlap",
        "semantic_rule_label",
        "semantic_rule_decision",
        "semantic_rule_id",
        "semantic_rule_reason",
        "static_validation_result",
        "evidence_party_role",
        "evidence_party_source_field",
        "evidence_artifact_url",
        "evidence_artifact_match",
        "evidence_artifact_match_method",
        "evidence_hash_usable_for_auto_repair",
        "observed_distribution_url",
        "gt_distribution_url",
        "hash_artifact_alignment",
        "hash_comparison_mode",
        "primary_evaluable",
        "primary_exclusion_reason",
        "project_consistent",
        "name_consistent",
        "proposal_component_id",
        "proposal_project_id",
        "proposal_tool",
        "proposal_sbom_path",
        "proposal_component_index",
        "proposal_row_number",
        "duplicate_proposal_rows_for_key",
    ]
    write_csv(
        output_dir / "gt_semantic_comparison_details.csv",
        matched,
        detail_fields,
    )

    unmatched_rows = [
        {
            "gt_id": case["gt_id"],
            "field_name": case["field_name"],
            "ecosystem": case["ecosystem"],
            "project_id": case["project_id"],
            "name": case["name"],
            "source_sbom": case["source_sbom"],
            "tool": case["gt_tool"],
            "component_index": case["component_index"],
            "version": case["version_for_identity"],
            "purl": case["purl_for_identity"],
            "ground_truth_status": case["ground_truth_status"],
            "ground_truth_value": case["ground_truth_value"],
            "reason": "no_matching_semantic_proposal",
        }
        for case in unmatched
    ]
    write_csv(output_dir / "verified_gt_cases_without_proposal.csv", unmatched_rows)

    write_csv(output_dir / "decision_vs_verified_gt_summary.csv", decisions)
    write_csv(
        output_dir / "decision_vs_verified_gt_by_field.csv",
        decision_summary_by_field(matched, unmatched),
    )
    write_csv(
        output_dir / "decision_value_match_matrix_all_fields.csv",
        decision_value_matrix(matched),
    )
    write_csv(
        output_dir / "decision_value_match_matrix_primary.csv",
        decision_value_matrix(primary_records(matched)),
    )
    write_csv(
        output_dir / "strict_precision_by_field.csv",
        precision_by_field(matched, {"accept"}),
    )
    write_csv(
        output_dir / "field_recovery_summary.csv",
        field_recovery_summary(gt_cases, matched, args.include_hash_primary),
    )
    write_csv(
        output_dir / "binary_acceptance_confusion_diagnostics.csv",
        [strict_confusion, expanded_confusion],
    )

    # Supplier-specific output is useful because supplier is intentionally
    # abstention-heavy in the revised SHM policy.
    supplier_details = [
        row for row in matched if row["field_name"] == "supplier"
    ]
    write_csv(
        output_dir / "supplier_verified_gt_comparison_details.csv",
        supplier_details,
        detail_fields,
    )

    # Hash diagnostic output remains separate until exact-artifact GT is cleaned.
    hash_details = [
        row for row in matched if row["field_name"] == "hash_values"
    ]
    write_csv(
        output_dir / "hash_verified_gt_diagnostic.csv",
        hash_details,
        detail_fields,
    )

    overall_summary = [
        {"metric": "ground_truth_component_occurrences", "value": len(gt_rows)},
        {"metric": "originally_missing_target_field_cases", "value": all_missing},
        {"metric": "verified_supported_gt_missing_cases_all_fields", "value": len(gt_cases)},
        {"metric": "gt_cases_with_semantic_proposal_all_fields", "value": len(matched)},
        {"metric": "gt_cases_without_semantic_proposal_all_fields", "value": len(unmatched)},
        {"metric": "proposal_coverage_all_fields_pct", "value": f"{pct(len(matched), len(gt_cases)):.2f}"},
        {"metric": "primary_hash_included", "value": str(args.include_hash_primary).lower()},
        {"metric": "primary_verified_gt_cases", "value": e2e["primary_verified_gt_cases"]},
        {"metric": "primary_gt_cases_with_proposal", "value": e2e["primary_gt_cases_with_proposal"]},
        {"metric": "primary_proposal_coverage_pct", "value": e2e["primary_proposal_coverage_pct"]},
        {"metric": "strict_accepted_evaluated", "value": strict["evaluated"]},
        {"metric": "strict_correct_accepted", "value": strict["correct"]},
        {"metric": "strict_incorrect_accepted", "value": strict["incorrect"]},
        {"metric": "strict_precision_pct", "value": f"{strict['precision_pct']:.2f}"},
        {"metric": "verified_repair_recall_pct", "value": e2e["verified_repair_recall_pct"]},
        {"metric": "end_to_end_f1_pct", "value": e2e["end_to_end_f1_pct"]},
        {"metric": "accept_plus_caution_evaluated", "value": expanded["evaluated"]},
        {"metric": "accept_plus_caution_correct", "value": expanded["correct"]},
        {"metric": "accept_plus_caution_incorrect", "value": expanded["incorrect"]},
        {"metric": "accept_plus_caution_precision_pct", "value": f"{expanded['precision_pct']:.2f}"},
    ]
    write_csv(output_dir / "ground_truth_evaluation_summary.csv", overall_summary)

    print("\nSaved outputs:")
    for name in [
        "ground_truth_evaluation_summary.csv",
        "decision_vs_verified_gt_summary.csv",
        "decision_vs_verified_gt_by_field.csv",
        "gt_semantic_comparison_details.csv",
        "strict_precision_by_field.csv",
        "field_recovery_summary.csv",
        "binary_acceptance_confusion_diagnostics.csv",
        "supplier_verified_gt_comparison_details.csv",
        "hash_verified_gt_diagnostic.csv",
        "verified_gt_cases_without_proposal.csv",
    ]:
        print(f"  {output_dir / name}")

    print("\nIMPORTANT INTERPRETATION NOTES")
    print("------------------------------")
    print("1. Strict precision uses semantic_rule_decision='accept' only.")
    print("2. accept_with_caution is reported separately from the primary result.")
    print("3. manual_review is an abstention, not automatically an error.")
    print("4. TP/FP/FN/TN is a secondary candidate-decision diagnostic only.")
    print("5. End-to-end TN/specificity is not defined because the GT set contains")
    print("   verified repair opportunities rather than a representative negative set.")
    if not args.include_hash_primary:
        print("6. hash_values enters primary metrics only when the original SBOM")
        print("   uniquely identifies the artifact; ambiguous hash cases remain diagnostic.")
        print("7. For an exact artifact, GT hashes are matched as a subset of SHM hashes;")
        print("   extra valid algorithms for that same artifact do not make the repair wrong.")


if __name__ == "__main__":
    main()
